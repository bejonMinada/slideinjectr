from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.schemas.presentation import DataFileInfo, DataSheetInfo

CSV_SHEET_NAME = "CSV_DATA"


def _csv_columns(path: Path) -> list[str]:
    df = pd.read_csv(path, nrows=0)
    return [str(col) for col in df.columns]


def _excel_sheet_info(path: Path) -> list[DataSheetInfo]:
    sheet_map = pd.read_excel(path, sheet_name=None)
    workbook = load_workbook(path, data_only=True, read_only=False)

    sheet_infos: list[DataSheetInfo] = []
    for sheet_name, df in sheet_map.items():
        ws = workbook[sheet_name]
        table_names = list(ws.tables.keys()) if ws.tables else []
        columns = [str(col) for col in df.columns]
        sheet_infos.append(
            DataSheetInfo(
                sheet_name=sheet_name,
                columns=columns,
                row_count=int(len(df.index)),
                table_names=table_names,
            )
        )

    workbook.close()
    return sheet_infos


def analyze_data_file(file_name: str, path: Path) -> DataFileInfo:
    suffix = path.suffix.lower()

    if suffix == ".csv":
        columns = _csv_columns(path)
        row_count = int(len(pd.read_csv(path).index))
        return DataFileInfo(
            file_name=file_name,
            file_type="csv",
            sheets=[
                DataSheetInfo(
                    sheet_name=CSV_SHEET_NAME,
                    columns=columns,
                    row_count=row_count,
                    table_names=[],
                )
            ],
        )

    file_type = "xlsx" if suffix == ".xlsx" else "xls"
    return DataFileInfo(file_name=file_name, file_type=file_type, sheets=_excel_sheet_info(path))
